from pathlib import Path
from fastapi import UploadFile, HTTPException
from pypdf import PdfReader
from docx import Document
from backend.utils.config import settings, UPLOAD_DIR
import uuid

ALLOWED = {".pdf", ".docx"}


def validate_upload(file: UploadFile, content: bytes):
    suffix = Path(file.filename or "").suffix.lower()
    if suffix not in ALLOWED:
        raise HTTPException(400, "Only PDF and DOCX files are allowed")
    if len(content) > settings.max_upload_mb * 1024 * 1024:
        raise HTTPException(413, f"File exceeds the {settings.max_upload_mb} MB limit")
    if not content:
        raise HTTPException(400, "Uploaded file is empty")
    return suffix


def extract_text(content: bytes, suffix: str) -> str:
    try:
        if suffix == ".pdf":
            from io import BytesIO

            reader = PdfReader(BytesIO(content))
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        else:
            from io import BytesIO

            doc = Document(BytesIO(content))
            text = "\n".join(p.text for p in doc.paragraphs)
            for table in doc.tables:
                for row in table.rows:
                    text += "\n" + " | ".join(cell.text for cell in row.cells)
    except Exception as exc:
        raise HTTPException(
            400, f"Could not parse the uploaded {suffix[1:].upper()} file"
        ) from exc
    text = " ".join(text.split())
    if len(text.strip()) < 20:
        raise HTTPException(400, "Resume contains too little readable text")
    return text


def save_file(content: bytes, suffix: str) -> str:
    filename = f"{uuid.uuid4().hex}{suffix}"
    (UPLOAD_DIR / filename).write_bytes(content)
    return filename
